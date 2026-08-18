"""Story 21.17: the server-side Excel export.

Golden-file style: every assertion opens the generated workbook with openpyxl and inspects
it, rather than trusting that the builder was called. Parity with the retired exceljs output
is the whole point, so the tests check sheet names, header rows in order, row content, the
hyperlink target, and the red font ARGB.

The column orders here are load-bearing — Story 8.23 deliberately places PyPI Latest
immediately before conda-forge Latest — so the header assertions compare full lists.
"""

from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path

import pytest
from django.core.files.base import ContentFile
from django.core.files.storage import default_storage
from django.test import Client
from openpyxl import load_workbook

from inventory.analysis.excel import (
    RED_ARGB,
    SheetSpec,
    build_workbook,
    licenses_sheet,
    safe_sheet_name,
    sbom_components_sheet,
    version_currency_sheet,
    vulnerabilities_sheet,
)
from inventory.analysis.models import AnalysisReport
from inventory.manifests.models import ManifestUpload
from inventory.sbom.models import SBOMJob
from inventory.users.models import Org
from inventory.users.services import create_org, register_user

PASSWORD = "pw12345678"

VERSION_REPORT = {
    "packages": [
        {
            "name": "divergent",
            "installed": "1.0.0",
            "latest": "2.0.0",
            "currency": "behind-1",
            "lts": "1.4",
            "on_lts": False,
            "ecosystem": "pypi",
            "conda_latest": "3.0.0",
            "latest_mismatch": True,
        },
        {
            "name": "aligned",
            "installed": "2.0.0",
            "latest": "2.0.0",
            "currency": "current",
            "lts": None,
            "on_lts": None,
            "ecosystem": None,
            "conda_latest": "2.0.0",
            "latest_mismatch": False,
        },
    ],
    "summary": {},
}

VULN_REPORT = {
    "packages": [
        {
            "name": "vulnpkg",
            "version": "1.0",
            "vulnerabilities": [
                {
                    "id": "GHSA-x",
                    "aliases": ["CVE-2024-1"],
                    "cvss_score": 9.1,
                    "severity": "Critical",
                    "cwe": ["CWE-79", "CWE-89"],
                    "advisory_url": "https://osv.dev/GHSA-x",
                }
            ],
        }
    ],
    "summary": {},
}

LICENSE_REPORT = {
    "tiers": [
        {"tier": "Strong Copyleft", "packages": [{"name": "gplpkg", "version": "1.0", "license": "GPL-3.0"}]},
        {"tier": "Permissive", "packages": [{"name": "mitpkg", "version": "2.0", "license": "MIT"}]},
    ],
    "summary": {},
}

SBOM_DOC = {
    "bomFormat": "CycloneDX",
    "specVersion": "1.6",
    "metadata": {"component": {"name": "app", "type": "application"}},
    "components": [
        {
            "type": "library",
            "name": "compa",
            "version": "1.0",
            "purl": "pkg:pypi/compa@1.0",
            "licenses": [{"license": {"id": "MIT"}}],
            "properties": [
                {"name": "package:ecosystem", "value": "pypi"},
                {"name": "sbom:relationship", "value": "direct"},
            ],
        }
    ],
}


def _sheet(workbook_bytes: bytes, name: str):  # type: ignore[no-untyped-def]
    return load_workbook(BytesIO(workbook_bytes))[name]


def _rows(sheet) -> list[list[object]]:  # type: ignore[no-untyped-def]
    """Read a sheet's cells, normalising blanks to empty strings.

    openpyxl stores an empty string as a *blank* cell and reads it back as None. That is a
    round-trip artifact, not a difference the user can see — a spreadsheet renders both as an
    empty cell — so normalising keeps the assertions about content rather than storage.
    """
    return [[("" if cell.value is None else cell.value) for cell in row] for row in sheet.iter_rows()]


# --- AC #2: the four builders, checked against the retired exceljs specs -------------------


def test_version_currency_sheet_columns_and_rows() -> None:
    workbook = build_workbook([version_currency_sheet(VERSION_REPORT["packages"])])
    sheet = _sheet(workbook, "Version Currency")
    rows = _rows(sheet)

    assert rows[0] == ["Package", "Installed", "PyPI Latest", "conda-forge Latest", "Status", "LTS", "On LTS", "Source"]
    assert rows[1] == ["divergent", "1.0.0", "2.0.0", "3.0.0", "behind-1", "1.4", "no", "pypi"]
    # on_lts None renders empty, not "no" — unknown is not the same as "not on LTS".
    assert rows[2] == ["aligned", "2.0.0", "2.0.0", "2.0.0", "current", "", "", ""]


def test_pypi_latest_immediately_precedes_conda_forge_latest() -> None:
    # Story 8.23 put them adjacent so they can be compared; the Dev Notes say not to "tidy" it.
    headers = _rows(_sheet(build_workbook([version_currency_sheet(VERSION_REPORT["packages"])]), "Version Currency"))[0]
    assert headers.index("conda-forge Latest") == headers.index("PyPI Latest") + 1


def test_vulnerabilities_sheet_columns_and_rows() -> None:
    sheet = _sheet(build_workbook([vulnerabilities_sheet(VULN_REPORT)]), "Vulnerabilities")
    rows = _rows(sheet)

    assert rows[0] == ["Package", "Installed", "CVE / GHSA", "CVSS", "Severity", "CWE", "Advisory URL"]
    # ids and cwe are comma-joined exactly as the SPA rendered them.
    assert rows[1] == [
        "vulnpkg",
        "1.0",
        "GHSA-x, CVE-2024-1",
        9.1,
        "Critical",
        "CWE-79, CWE-89",
        "https://osv.dev/GHSA-x",
    ]


def test_licenses_sheet_flattens_every_tier() -> None:
    sheet = _sheet(build_workbook([licenses_sheet(LICENSE_REPORT)]), "Licenses")
    rows = _rows(sheet)

    assert rows[0] == ["Package", "Installed", "License", "Risk Tier"]
    # The tier travels with the row so a flat sheet keeps the risk grouping.
    assert rows[1] == ["gplpkg", "1.0", "GPL-3.0", "Strong Copyleft"]
    assert rows[2] == ["mitpkg", "2.0", "MIT", "Permissive"]


def test_sbom_components_sheet_adds_optional_columns_only_when_present() -> None:
    from inventory.sbom.document import normalize_components

    components = normalize_components(json.dumps(SBOM_DOC).encode(), "cyclonedx-json")
    sheet = _sheet(build_workbook([sbom_components_sheet(components)]), "SBOM Components")
    rows = _rows(sheet)

    assert rows[0][:5] == ["Name", "Version", "Type", "License", "Relationship"]
    assert "Ecosystem" in rows[0] and "PURL" in rows[0]


def test_sbom_components_sheet_omits_purl_when_no_component_has_one() -> None:
    sheet = _sheet(build_workbook([sbom_components_sheet([{"name": "x", "version": "1"}])]), "SBOM Components")
    headers = _rows(sheet)[0]
    assert "PURL" not in headers
    assert "Ecosystem" not in headers


# --- AC #3: styling parity, verified in the workbook --------------------------------------


def test_the_divergence_cell_carries_the_theme_red() -> None:
    """Story 8.22's red must survive into the sheet, not just the page."""
    sheet = _sheet(build_workbook([version_currency_sheet(VERSION_REPORT["packages"])]), "Version Currency")

    diverging = sheet.cell(row=2, column=4)  # conda-forge Latest, divergent package
    aligned = sheet.cell(row=3, column=4)

    assert diverging.value == "3.0.0"
    assert diverging.font.color.rgb == RED_ARGB
    # And the matching package is NOT red — otherwise the colour would mean nothing.
    assert aligned.font.color is None or aligned.font.color.rgb != RED_ARGB


def test_the_package_name_is_a_working_hyperlink() -> None:
    sheet = _sheet(build_workbook([version_currency_sheet(VERSION_REPORT["packages"])]), "Version Currency")

    linked = sheet.cell(row=2, column=1)
    unlinked = sheet.cell(row=3, column=1)  # ecosystem None → no registry URL

    assert linked.hyperlink is not None
    assert linked.hyperlink.target == "https://pypi.org/project/divergent/1.0.0/"
    assert unlinked.hyperlink is None


def test_header_rows_are_bold() -> None:
    sheet = _sheet(build_workbook([licenses_sheet(LICENSE_REPORT)]), "Licenses")
    assert all(cell.font.bold for cell in sheet[1])


# --- AC #6: sheet names stay Excel-legal --------------------------------------------------


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("Licenses", "Licenses"),
        ("a/b:c*d?e[f]g\\h", "a-b-c-d-e-f-g-h"),
        ("x" * 40, "x" * 31),
        ("", "Sheet"),
    ],
)
def test_sheet_names_are_sanitised(raw: str, expected: str) -> None:
    # Enforced by the service, not trusted from the caller: openpyxl raises on an illegal
    # name, which would turn a download into a 500.
    assert safe_sheet_name(raw) == expected


def test_an_illegal_sheet_name_still_produces_a_workbook() -> None:
    workbook = build_workbook([SheetSpec(name="bad/name", columns=[("a", "A")], rows=[{"a": 1}])])
    assert "bad-name" in load_workbook(BytesIO(workbook)).sheetnames


# --- The export endpoints -------------------------------------------------------------------


def _client(email: str) -> Client:
    client = Client()
    assert client.login(email=email, password=PASSWORD)
    return client


def _job(org: Org, *, reports: dict[str, dict | None] | None = None, with_sbom: bool = True) -> SBOMJob:
    upload = ManifestUpload.objects.create(
        org=org,
        file="m/f.txt",
        detected_format=ManifestUpload.Format.REQUIREMENTS,
        original_filename="requirements.txt",
        application_id="APP",
        component_name="c",
        repository_url="https://example.com/r",
        source_branch="main",
    )
    result_key = None
    if with_sbom:
        result_key = f"sboms/{upload.pk}.json"
        default_storage.save(result_key, ContentFile(json.dumps(SBOM_DOC)))
    job = SBOMJob.objects.create(
        org=org,
        manifest=upload,
        output_format="cyclonedx-json",
        status=SBOMJob.Status.SUCCESS,
        result_key=result_key,
        summary_stats={"total_packages": 2},
    )
    for report_type, payload in (reports or {}).items():
        key = None
        if payload is not None:
            key = f"reports/{job.task_id}-{report_type}.json"
            default_storage.save(key, ContentFile(json.dumps(payload)))
        AnalysisReport.objects.create(
            job=job,
            report_type=report_type,
            artifact_key=key,
            failed=payload is None,
            failure_reason="phase_failed" if payload is None else None,
        )
    return job


ALL_REPORTS = {
    AnalysisReport.ReportType.VULN: VULN_REPORT,
    AnalysisReport.ReportType.LICENSE: LICENSE_REPORT,
    AnalysisReport.ReportType.VERSION: VERSION_REPORT,
}


@pytest.fixture
def org_client():  # type: ignore[no-untyped-def]
    user = register_user(email="dev@example.com", password=PASSWORD)
    org = create_org(name="Acme", admin_user=user)
    return _client("dev@example.com"), org


@pytest.mark.django_db
@pytest.mark.parametrize(
    ("kind", "sheet_name"),
    [
        ("vulnerabilities", "Vulnerabilities"),
        ("licenses", "Licenses"),
        ("versions", "Version Currency"),
        ("sbom", "SBOM Components"),
    ],
)
def test_each_tab_exports_its_own_sheet(org_client, kind: str, sheet_name: str) -> None:  # type: ignore[no-untyped-def]
    client, org = org_client
    job = _job(org, reports=ALL_REPORTS)

    response = client.get(f"/results/{job.task_id}/export/{kind}.xlsx")

    assert response.status_code == 200
    assert response["Content-Type"].endswith("spreadsheetml.sheet")
    assert "attachment;" in response["Content-Disposition"]
    assert load_workbook(BytesIO(response.content)).sheetnames == [sheet_name]


# --- AC #4: the combined workbook ------------------------------------------------------------


@pytest.mark.django_db
def test_the_combined_workbook_contains_every_available_sheet(org_client) -> None:  # type: ignore[no-untyped-def]
    client, org = org_client
    job = _job(org, reports=ALL_REPORTS)

    response = client.get(f"/results/{job.task_id}/export.xlsx")

    assert load_workbook(BytesIO(response.content)).sheetnames == [
        "SBOM Components",
        "Vulnerabilities",
        "Licenses",
        "Version Currency",
    ]


@pytest.mark.django_db
def test_a_failed_phase_is_omitted_rather_than_emitted_empty(org_client) -> None:  # type: ignore[no-untyped-def]
    """An empty sheet would read as "we checked and found nothing", which is a false claim."""
    client, org = org_client
    job = _job(org, reports={**ALL_REPORTS, AnalysisReport.ReportType.VULN: None})

    response = client.get(f"/results/{job.task_id}/export.xlsx")

    names = load_workbook(BytesIO(response.content)).sheetnames
    assert "Vulnerabilities" not in names
    assert "Licenses" in names and "Version Currency" in names


@pytest.mark.django_db
def test_exporting_a_failed_report_on_its_own_is_404(org_client) -> None:  # type: ignore[no-untyped-def]
    client, org = org_client
    job = _job(org, reports={AnalysisReport.ReportType.VULN: None})

    assert client.get(f"/results/{job.task_id}/export/vulnerabilities.xlsx").status_code == 404


@pytest.mark.django_db
def test_a_job_with_nothing_to_export_is_404(org_client) -> None:  # type: ignore[no-untyped-def]
    client, org = org_client
    job = _job(org, with_sbom=False)

    assert client.get(f"/results/{job.task_id}/export.xlsx").status_code == 404


# --- AC #5: org scoping -----------------------------------------------------------------------


@pytest.mark.django_db
def test_exports_are_org_scoped_and_indistinguishable_from_missing(org_client) -> None:  # type: ignore[no-untyped-def]
    client, _ = org_client
    outsider = register_user(email="outsider@example.com", password=PASSWORD)
    other_org = create_org(name="Other", admin_user=outsider)
    theirs = _job(other_org, reports=ALL_REPORTS)

    cross_org = client.get(f"/results/{theirs.task_id}/export.xlsx")
    missing = client.get("/results/00000000-0000-0000-0000-000000000000/export.xlsx")

    assert cross_org.status_code == missing.status_code == 404


@pytest.mark.django_db
def test_anonymous_cannot_export(org_client) -> None:  # type: ignore[no-untyped-def]
    _, org = org_client
    job = _job(org, reports=ALL_REPORTS)

    response = Client().get(f"/results/{job.task_id}/export.xlsx")

    assert response.status_code == 302
    assert response.headers["Location"].startswith("/login")


@pytest.mark.django_db
def test_an_unknown_export_kind_is_404(org_client) -> None:  # type: ignore[no-untyped-def]
    client, org = org_client
    job = _job(org, reports=ALL_REPORTS)
    assert client.get(f"/results/{job.task_id}/export/wat.xlsx").status_code == 404


@pytest.mark.django_db
def test_the_tabs_link_to_their_exports(org_client) -> None:  # type: ignore[no-untyped-def]
    client, org = org_client
    job = _job(org, reports=ALL_REPORTS)

    overview = client.get(f"/results/{job.task_id}").content.decode()
    assert f"/results/{job.task_id}/export.xlsx" in overview

    for kind in ("vulnerabilities", "licenses", "versions", "sbom"):
        tab = client.get(f"/results/{job.task_id}/tab/{kind}").content.decode()
        assert f"/results/{job.task_id}/export/{kind}.xlsx" in tab


# --- AC #3/Task 5: parity against a workbook the RETIRED exceljs pipeline produced ---------


REFERENCE_WORKBOOK = (
    Path(__file__).resolve().parents[1] / "fixtures" / "excel_reference" / "version_currency_exceljs.xlsx"
)


def test_the_reference_workbook_is_committed() -> None:
    """Story 21.19 deletes `frontend/`, taking exceljs with it.

    The fixture was generated from the real exceljs pipeline while it still ran, so parity
    remains checkable afterwards. Once it is gone there is no way to regenerate it, which is
    why the Dev Notes insisted on capturing it during this story — without it AC #3 is an
    assertion of faith.
    """
    assert REFERENCE_WORKBOOK.is_file(), "the exceljs reference workbook is missing and cannot be regenerated"


def test_output_matches_the_exceljs_reference_cell_for_cell() -> None:
    reference = load_workbook(REFERENCE_WORKBOOK)
    ours = load_workbook(BytesIO(build_workbook([version_currency_sheet(VERSION_REPORT["packages"])])))

    assert ours.sheetnames == reference.sheetnames

    theirs_sheet = reference["Version Currency"]
    our_sheet = ours["Version Currency"]
    assert _rows(our_sheet) == _rows(theirs_sheet)


def test_styling_matches_the_exceljs_reference() -> None:
    """The part a cell-value diff cannot see: hyperlink targets and font colour."""
    reference = load_workbook(REFERENCE_WORKBOOK)["Version Currency"]
    ours = load_workbook(BytesIO(build_workbook([version_currency_sheet(VERSION_REPORT["packages"])])))[
        "Version Currency"
    ]

    # The divergence cell is red in both, with the same ARGB.
    assert ours.cell(row=2, column=4).font.color.rgb == reference.cell(row=2, column=4).font.color.rgb == RED_ARGB
    # The package name links to the same target in both.
    assert ours.cell(row=2, column=1).hyperlink.target == reference.cell(row=2, column=1).hyperlink.target
    # Headers are bold in both.
    assert [cell.font.bold for cell in ours[1]] == [cell.font.bold for cell in reference[1]]
